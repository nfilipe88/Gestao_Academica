"""Importação de Dados Legados (mini-pauta MININED) — ver app/cruds/importacao.py.

Os ficheiros de teste reproduzem a estrutura real confirmada em dois
ficheiros MININED genuínos (linha 12 = metadados, 13 = bandas de
trimestre, 14 = cabeçalhos, 15+ = alunos): a fórmula do MT do 1º
trimestre é uma média simples (SUM/2 -> 50/50), a do 2º/3º é ponderada
(0.4/0.6) — ver _construir_mini_pauta_xlsx.
"""
import io
from decimal import Decimal

from openpyxl import Workbook

from tests.conftest import auth_headers, criar_escola_e_gestor, sufixo_unico
from tests.test_comportamento import _criar_professor_com_token

_CONTENT_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_LINHA_METADADOS = 12
_LINHA_BANDAS = 13
_LINHA_CABECALHOS = 14
_LINHA_PRIMEIRO_ALUNO = 15


def _construir_mini_pauta_xlsx(
    *,
    disciplina: str = "Educação Laboral",
    sala: str = "04",
    turma: str = "AD7T/4",
    turno: str = "TARDE",
    ano_letivo_raw: str = "2025/2026",
    alunos: list[dict] | None = None,
    valor_solto: tuple[int, int, object] | None = None,
    sem_cabecalhos: bool = False,
) -> bytes:
    """alunos: lista de {"nome": str, "t1": (mact, pt) | None, ...}. Um
    valor None num trimestre = célula em branco nesse par MACT/PT."""
    if alunos is None:
        alunos = [{"nome": "Aluno Um", "t1": (15, 13)}, {"nome": "Aluno Dois", "t1": (10, 16)}]

    wb = Workbook()
    ws = wb.active

    if sem_cabecalhos:
        ws.cell(1, 1).value = "Ficheiro sem a estrutura esperada"
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    ws.cell(_LINHA_METADADOS, 1).value = f"     DISCIPLINA: {disciplina}"
    ws.cell(_LINHA_METADADOS, 3).value = f"SALA:   {sala}"
    ws.cell(_LINHA_METADADOS, 7).value = f"TURMA:   {turma}"
    ws.cell(_LINHA_METADADOS, 11).value = f"TURNO:    {turno}"
    ws.cell(_LINHA_METADADOS, 15).value = f"ANO LECTIVO:   {ano_letivo_raw}"

    ws.cell(_LINHA_BANDAS, 3).value = "1º TRIMESTRE"
    ws.cell(_LINHA_BANDAS, 7).value = "2º TRIMESTRE"
    ws.cell(_LINHA_BANDAS, 11).value = "3º TRIMESTRE"
    ws.cell(_LINHA_BANDAS, 15).value = "CLASSIFICAÇÃO GERAL"

    cabecalhos = ["Nº", "NOME DO ALUNO", "MACT", "PT", "MT", "F T", "MACT", "PT", "MT", "F T", "MACT", "PT", "MT", "F T", "MFD", "MEO", "NEN", "M. FINAL"]
    for col, texto in enumerate(cabecalhos, start=1):
        ws.cell(_LINHA_CABECALHOS, col).value = texto

    for i, aluno in enumerate(alunos):
        r = _LINHA_PRIMEIRO_ALUNO + i
        ws.cell(r, 1).value = i + 1
        ws.cell(r, 2).value = aluno["nome"]
        t1 = aluno.get("t1")
        if t1:
            ws.cell(r, 3).value, ws.cell(r, 4).value = t1
        ws.cell(r, 5).value = f"=SUM(C{r}+D{r})/2"
        t2 = aluno.get("t2")
        if t2:
            ws.cell(r, 7).value, ws.cell(r, 8).value = t2
        ws.cell(r, 9).value = f'=IF(AND(G{r}="",H{r}=""),"",ROUND((0.4*G{r}+0.6*H{r}),0))'
        t3 = aluno.get("t3")
        if t3:
            ws.cell(r, 11).value, ws.cell(r, 12).value = t3
        ws.cell(r, 13).value = f'=IF(AND(K{r}="",L{r}=""),"",ROUND((0.4*K{r}+0.6*L{r}),0))'

    # Linha de assinatura — sem Nº, marca o fim da lista de alunos (mesmo padrão dos ficheiros reais).
    linha_assinatura = _LINHA_PRIMEIRO_ALUNO + len(alunos)
    ws.cell(linha_assinatura, 2).value = "ASSINATURA DO PROFESSOR"

    if valor_solto:
        r, c, v = valor_solto
        ws.cell(r, c).value = v

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def _definir_nota_maxima(client, headers, valor: int = 20):
    """A mini-pauta MININED usa a escala 0-20 (MACT/PT podem passar de
    10) — Tenant.nota_maxima nasce em 10 (valor por omissão preservado
    para escolas já existentes, ver migração), por isso os testes que
    confirmam uma importação com notas > 10 têm de a alargar primeiro,
    tal como uma escola real faria em Configurações antes de importar."""
    resp = await client.get("/api/v1/configuracoes", headers=headers)
    assert resp.status_code == 200, resp.text
    config = resp.json()
    config["nota_maxima"] = valor
    resp = await client.put("/api/v1/configuracoes", headers=headers, json=config)
    assert resp.status_code == 200, resp.text


async def _preview(client, headers, conteudo: bytes, nome: str = "mini-pauta.xlsx"):
    return await client.post(
        "/api/v1/importacao/mini-pauta/preview", headers=headers,
        files={"ficheiro": (nome, io.BytesIO(conteudo), _CONTENT_TYPE_XLSX)}
    )


def _payload_confirmacao(preview_json: dict, *, curso="Ensino Geral", serie="7ª Classe", acoes: dict | None = None) -> dict:
    """Monta o corpo de /confirmar a partir de um preview, com CRIAR_NOVO_ALUNO
    por omissão para todos (acoes permite sobrepor nome->ação/aluno_existente_id)."""
    acoes = acoes or {}
    meta = preview_json["metadados"]
    alunos = []
    for a in preview_json["alunos"]:
        override = acoes.get(a["nome_completo"], {})
        alunos.append({
            "nome_completo": a["nome_completo"],
            "acao": override.get("acao", "CRIAR_NOVO_ALUNO"),
            "aluno_existente_id": override.get("aluno_existente_id", a["aluno_existente_id"]),
            "notas": a["notas"],
        })
    return {
        "curso": {"nome_novo": curso},
        "serie_ano": {"nome_novo": serie},
        "turma": {"nome_novo": meta["turma_nome_codigo"]},
        "disciplina": {"nome_novo": meta["disciplina_nome"]},
        "ano_letivo": meta["ano_letivo"],
        "trimestres": [
            {"nome": t["nome"], "tem_dados": t["tem_dados"], "peso_mact": t["peso_mact"], "peso_pt": t["peso_pt"]}
            for t in preview_json["trimestres"]
        ],
        "alunos": alunos,
        "nome_ficheiro_original": "mini-pauta.xlsx",
        "chave_storage_ficheiro": meta["chave_storage_ficheiro"],
    }


# ==========================================
# CAMINHO FELIZ
# ==========================================
async def test_importar_mini_pauta_caminho_feliz(client):
    escola = await criar_escola_e_gestor(client, "importacao-feliz")
    headers = auth_headers(escola["token"])
    await _definir_nota_maxima(client, headers)
    conteudo = _construir_mini_pauta_xlsx(alunos=[
        {"nome": "Fulano da Silva", "t1": (15, 13)},
        {"nome": "Fulana Exemplo", "t1": (10, 16)},
        {"nome": "Beltrano Sousa", "t1": (7, 9)},
    ])

    resp = await _preview(client, headers, conteudo)
    assert resp.status_code == 200, resp.text
    preview = resp.json()

    assert preview["metadados"]["disciplina_nome"] == "Educação Laboral"
    assert preview["metadados"]["turma_nome_codigo"] == "AD7T/4"
    assert preview["metadados"]["ano_letivo"] == 2025
    assert len(preview["alunos"]) == 3

    t1, t2, t3 = preview["trimestres"]
    assert t1["tem_dados"] is True
    assert Decimal(str(t1["peso_mact"])) == Decimal("50") and Decimal(str(t1["peso_pt"])) == Decimal("50")
    assert t1["peso_detectado_automaticamente"] is True
    assert t2["tem_dados"] is False and t3["tem_dados"] is False
    # Sem dados no 2º/3º trimestre, mas o peso da fórmula continua reconhecido (40/60).
    assert Decimal(str(t2["peso_mact"])) == Decimal("40") and Decimal(str(t2["peso_pt"])) == Decimal("60")

    corpo = _payload_confirmacao(preview)
    resp = await client.post("/api/v1/importacao/mini-pauta/confirmar", headers=headers, json=corpo)
    assert resp.status_code == 201, resp.text
    resultado = resp.json()
    assert resultado["resumo"]["total_alunos_criados"] == 3
    assert resultado["resumo"]["total_alunos_reaproveitados"] == 0
    # 1º trimestre tem dados -> 2 avaliações (MACT/PT) x 3 alunos = 6 notas. 2º/3º sem dados -> nada.
    assert resultado["resumo"]["total_notas_lancadas"] == 6
    assert set(resultado["alunos_com_data_nascimento_pendente"]) == {"Fulano da Silva", "Fulana Exemplo", "Beltrano Sousa"}

    resp = await client.get("/api/v1/academico/turmas", headers=headers)
    turmas = [t for t in resp.json() if t["nome_codigo"] == "AD7T/4"]
    assert len(turmas) == 1
    turma_id = turmas[0]["id"]
    assert turmas[0]["ano_letivo"] == 2025

    resp = await client.get("/api/v1/academico/disciplinas", headers=headers)
    disciplinas = [d for d in resp.json() if d["nome"] == "Educação Laboral"]
    assert len(disciplinas) == 1
    disciplina_id = disciplinas[0]["id"]

    resp = await client.get("/api/v1/alunos", headers=headers)
    nomes_alunos = {a["nome_completo"] for a in resp.json()["items"]}
    assert {"Fulano da Silva", "Fulana Exemplo", "Beltrano Sousa"} <= nomes_alunos

    resp = await client.get(
        f"/api/v1/diario/turmas/{turma_id}/disciplinas/{disciplina_id}/notas-finais",
        headers=headers, params={"periodo_avaliacao": "1º TRIMESTRE"}
    )
    assert resp.status_code == 200, resp.text
    notas_finais = {n["nome_aluno"]: n["valor_nota"] for n in resp.json()}
    # MT = média simples 50/50 -> (MACT+PT)/2, arredondado a 2 casas pelo motor de cálculo.
    assert notas_finais["Fulano da Silva"] == 14.0
    assert notas_finais["Fulana Exemplo"] == 13.0
    assert notas_finais["Beltrano Sousa"] == 8.0
    assert all(n["calculada_automaticamente"] for n in (await client.get(
        f"/api/v1/diario/turmas/{turma_id}/disciplinas/{disciplina_id}/notas-finais",
        headers=headers, params={"periodo_avaliacao": "1º TRIMESTRE"}
    )).json())


# ==========================================
# RBAC
# ==========================================
async def test_professor_nao_pode_importar(client):
    escola = await criar_escola_e_gestor(client, "importacao-rbac")
    headers = auth_headers(escola["token"])
    _, token_professor = await _criar_professor_com_token(client, headers, "Prof. Importação")
    headers_professor = auth_headers(token_professor)

    conteudo = _construir_mini_pauta_xlsx()
    resp = await _preview(client, headers_professor, conteudo)
    assert resp.status_code == 403, resp.text

    resp = await client.post(
        "/api/v1/importacao/mini-pauta/confirmar", headers=headers_professor,
        json={
            "curso": {"nome_novo": "X"}, "serie_ano": {"nome_novo": "Y"}, "turma": {"nome_novo": "Z"},
            "disciplina": {"nome_novo": "W"}, "ano_letivo": 2025, "trimestres": [], "alunos": [],
            "nome_ficheiro_original": "x.xlsx", "chave_storage_ficheiro": "x",
        }
    )
    assert resp.status_code == 403, resp.text


# ==========================================
# FICHEIRO MAL FORMADO
# ==========================================
async def test_ficheiro_sem_estrutura_esperada_e_rejeitado(client):
    escola = await criar_escola_e_gestor(client, "importacao-malformado")
    headers = auth_headers(escola["token"])
    conteudo = _construir_mini_pauta_xlsx(sem_cabecalhos=True)

    resp = await _preview(client, headers, conteudo)
    assert resp.status_code == 400, resp.text

    resp = await client.get("/api/v1/alunos", headers=headers)
    assert resp.json()["items"] == []


async def test_tipo_de_ficheiro_invalido_e_rejeitado(client):
    escola = await criar_escola_e_gestor(client, "importacao-tipo-invalido")
    headers = auth_headers(escola["token"])
    resp = await client.post(
        "/api/v1/importacao/mini-pauta/preview", headers=headers,
        files={"ficheiro": ("nota.txt", io.BytesIO(b"nao e um xlsx"), "text/plain")}
    )
    assert resp.status_code == 400, resp.text


# ==========================================
# MACT+PT EM BRANCO
# ==========================================
async def test_mact_pt_em_branco_nao_lanca_nota_zero(client):
    escola = await criar_escola_e_gestor(client, "importacao-branco")
    headers = auth_headers(escola["token"])
    await _definir_nota_maxima(client, headers)
    conteudo = _construir_mini_pauta_xlsx(alunos=[
        {"nome": "Com Notas", "t1": (12, 10)},
        {"nome": "Sem Notas", "t1": None},
    ])

    resp = await _preview(client, headers, conteudo)
    preview = resp.json()
    sem_notas = next(a for a in preview["alunos"] if a["nome_completo"] == "Sem Notas")
    assert any("branco" in anomalia for anomalia in sem_notas["anomalias"])
    assert sem_notas["notas"]["1º TRIMESTRE"]["mact"] is None

    corpo = _payload_confirmacao(preview)
    resp = await client.post("/api/v1/importacao/mini-pauta/confirmar", headers=headers, json=corpo)
    assert resp.status_code == 201, resp.text
    # 1 aluno com nota -> 2 avaliações (MACT/PT) x 1 = 2 notas, nunca 4.
    assert resp.json()["resumo"]["total_notas_lancadas"] == 2

    resp = await client.get("/api/v1/academico/turmas", headers=headers)
    turma_id = next(t for t in resp.json() if t["nome_codigo"] == "AD7T/4")["id"]
    resp = await client.get("/api/v1/academico/disciplinas", headers=headers)
    disciplina_id = next(d for d in resp.json() if d["nome"] == "Educação Laboral")["id"]

    resp = await client.get(
        f"/api/v1/diario/turmas/{turma_id}/disciplinas/{disciplina_id}/notas-finais",
        headers=headers, params={"periodo_avaliacao": "1º TRIMESTRE"}
    )
    por_nome = {n["nome_aluno"]: n["valor_nota"] for n in resp.json()}
    assert por_nome["Com Notas"] == 11.0
    assert por_nome.get("Sem Notas") is None  # nunca 0 — sem linha nenhuma criada


# ==========================================
# PREVIEW NÃO GRAVA NADA
# ==========================================
async def test_preview_repetido_nao_cria_nada(client):
    escola = await criar_escola_e_gestor(client, "importacao-preview-idempotente")
    headers = auth_headers(escola["token"])
    conteudo = _construir_mini_pauta_xlsx()

    await _preview(client, headers, conteudo)
    await _preview(client, headers, conteudo)

    resp = await client.get("/api/v1/alunos", headers=headers)
    assert resp.json()["items"] == []
    resp = await client.get("/api/v1/academico/turmas", headers=headers)
    assert resp.json() == []
    resp = await client.get("/api/v1/importacao/lotes", headers=headers)
    assert resp.json() == []


# ==========================================
# NOME DUPLICADO — aviso, mas não bloqueia criar um segundo aluno distinto
# ==========================================
async def test_nome_duplicado_avisa_mas_permite_criar_novo(client):
    escola = await criar_escola_e_gestor(client, "importacao-duplicado")
    headers = auth_headers(escola["token"])

    resp = await client.post("/api/v1/alunos", headers=headers, json={
        "matricula_interna": f"AL{sufixo_unico()}", "nome_completo": "Nome Repetido", "data_nascimento": "2012-01-01"
    })
    assert resp.status_code == 201, resp.text
    aluno_existente_id = resp.json()["id"]

    conteudo = _construir_mini_pauta_xlsx(alunos=[{"nome": "Nome Repetido", "t1": (10, 10)}])
    resp = await _preview(client, headers, conteudo)
    preview = resp.json()
    aluno_detectado = preview["alunos"][0]
    assert aluno_detectado["aluno_existente_id"] == aluno_existente_id
    assert any("Já existe um aluno" in a for a in aluno_detectado["anomalias"])

    corpo = _payload_confirmacao(preview)  # ação por omissão = CRIAR_NOVO_ALUNO
    resp = await client.post("/api/v1/importacao/mini-pauta/confirmar", headers=headers, json=corpo)
    assert resp.status_code == 201, resp.text
    assert resp.json()["resumo"]["total_alunos_criados"] == 1

    resp = await client.get("/api/v1/alunos", headers=headers)
    repetidos = [a for a in resp.json()["items"] if a["nome_completo"] == "Nome Repetido"]
    assert len(repetidos) == 2  # o já existente + o novo criado pelo import


# ==========================================
# _extrair_pesos_formula
# ==========================================
def test_extrair_pesos_formula_media_simples():
    from app.cruds.importacao import _extrair_pesos_formula
    peso_mact, peso_pt, detectado = _extrair_pesos_formula("=SUM(C15+D15)/2", "C", "D")
    assert (peso_mact, peso_pt, detectado) == (Decimal("50"), Decimal("50"), True)


def test_extrair_pesos_formula_ponderada():
    from app.cruds.importacao import _extrair_pesos_formula
    formula = '=IF(AND(G15="",H15=""),"",ROUND((0.4*G15+0.6*H15),0))'
    peso_mact, peso_pt, detectado = _extrair_pesos_formula(formula, "G", "H")
    assert (peso_mact, peso_pt, detectado) == (Decimal("40"), Decimal("60"), True)


def test_extrair_pesos_formula_nao_reconhecida_cai_em_50_50():
    from app.cruds.importacao import _extrair_pesos_formula
    peso_mact, peso_pt, detectado = _extrair_pesos_formula("=MEDIA(C15:D15)", "C", "D")
    assert (peso_mact, peso_pt, detectado) == (Decimal("50"), Decimal("50"), False)


def test_extrair_pesos_formula_ausente_cai_em_50_50():
    from app.cruds.importacao import _extrair_pesos_formula
    peso_mact, peso_pt, detectado = _extrair_pesos_formula(None, "C", "D")
    assert (peso_mact, peso_pt, detectado) == (Decimal("50"), Decimal("50"), False)


# ==========================================
# VALOR SOLTO FORA DA TABELA
# ==========================================
async def test_valor_solto_fora_da_tabela_vira_so_anomalia(client):
    escola = await criar_escola_e_gestor(client, "importacao-valor-solto")
    headers = auth_headers(escola["token"])
    # Coluna 19 (S) fica fora do intervalo de cabeçalhos (1-18) — mesmo achado dos ficheiros reais.
    conteudo = _construir_mini_pauta_xlsx(
        alunos=[{"nome": "Aluno Alvo", "t1": (10, 10)}], valor_solto=(_LINHA_PRIMEIRO_ALUNO, 19, 2)
    )
    resp = await _preview(client, headers, conteudo)
    preview = resp.json()
    aluno = preview["alunos"][0]
    assert any("coluna S" in a and "fora da tabela" in a for a in aluno["anomalias"])
    # Nunca aparece como campo estruturado.
    assert all(k in ("1º TRIMESTRE", "2º TRIMESTRE", "3º TRIMESTRE") for k in aluno["notas"].keys())


# ==========================================
# DEDUP DE TURMA (duas disciplinas, mesma turma)
# ==========================================
async def test_duas_disciplinas_mesma_turma_nao_duplicam_turma(client):
    escola = await criar_escola_e_gestor(client, "importacao-dedup-turma")
    headers = auth_headers(escola["token"])

    conteudo_1 = _construir_mini_pauta_xlsx(disciplina="Educação Laboral", alunos=[{"nome": "Aluno A", "t1": (10, 10)}])
    resp = await _preview(client, headers, conteudo_1)
    corpo_1 = _payload_confirmacao(resp.json())
    resp = await client.post("/api/v1/importacao/mini-pauta/confirmar", headers=headers, json=corpo_1)
    assert resp.status_code == 201, resp.text

    conteudo_2 = _construir_mini_pauta_xlsx(disciplina="Matemática", alunos=[{"nome": "Aluno B", "t1": (10, 10)}])
    resp = await _preview(client, headers, conteudo_2)
    corpo_2 = _payload_confirmacao(resp.json())
    resp = await client.post("/api/v1/importacao/mini-pauta/confirmar", headers=headers, json=corpo_2)
    assert resp.status_code == 201, resp.text

    resp = await client.get("/api/v1/academico/turmas", headers=headers)
    turmas = [t for t in resp.json() if t["nome_codigo"] == "AD7T/4"]
    assert len(turmas) == 1  # find-or-create — não duplicou

    resp = await client.get("/api/v1/academico/disciplinas", headers=headers)
    nomes_disciplinas = {d["nome"] for d in resp.json()}
    assert {"Educação Laboral", "Matemática"} <= nomes_disciplinas


# ==========================================
# DESFAZER
# ==========================================
async def test_desfazer_importacao_remove_plantel_mas_mantem_estrutura(client):
    escola = await criar_escola_e_gestor(client, "importacao-desfazer")
    headers = auth_headers(escola["token"])
    conteudo = _construir_mini_pauta_xlsx(alunos=[{"nome": "Aluno Desfazer", "t1": (10, 10)}])

    resp = await _preview(client, headers, conteudo)
    corpo = _payload_confirmacao(resp.json())
    resp = await client.post("/api/v1/importacao/mini-pauta/confirmar", headers=headers, json=corpo)
    assert resp.status_code == 201, resp.text
    lote_id = resp.json()["lote_importacao_id"]

    resp = await client.post(f"/api/v1/importacao/lotes/{lote_id}/desfazer", headers=headers)
    assert resp.status_code == 200, resp.text
    resultado = resp.json()
    assert resultado["estado"] == "DESFEITO"
    assert resultado["alunos_removidos"] == 1
    assert resultado["alunos_mantidos"] == 0

    resp = await client.get("/api/v1/alunos", headers=headers)
    assert all(a["nome_completo"] != "Aluno Desfazer" for a in resp.json()["items"])

    # Turma/Disciplina (estrutura partilhada) continuam a existir.
    resp = await client.get("/api/v1/academico/turmas", headers=headers)
    assert any(t["nome_codigo"] == "AD7T/4" for t in resp.json())
    resp = await client.get("/api/v1/academico/disciplinas", headers=headers)
    assert any(d["nome"] == "Educação Laboral" for d in resp.json())

    resp = await client.get("/api/v1/importacao/lotes", headers=headers)
    lote = next(l for l in resp.json() if l["id"] == lote_id)
    assert lote["estado"] == "DESFEITO"

    # Desfazer duas vezes não é permitido.
    resp = await client.post(f"/api/v1/importacao/lotes/{lote_id}/desfazer", headers=headers)
    assert resp.status_code == 400, resp.text
